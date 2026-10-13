"""Build the Excel workbook (dashboard + data sheets) from the scraped DataFrame."""
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Reference categorical palette (fixed order: slot 1 blue, slot 2 orange).
SOURCE_COLORS = {"Confidence": "2A78D6", "JSK": "EB6834"}
INK, MUTED, RULE, TILE = "0B0B0B", "52514E", "D9D8D3", "F3F2EF"

SIZE_BANDS = [0, 1000, 1500, 2500, 5000, float("inf")]
SIZE_LABELS = ["< 1,000 m²", "1,000–1,500 m²", "1,500–2,500 m²", "2,500–5,000 m²", "5,000+ m²"]

# (column, header, width, number format)
DATA_COLUMNS = [
    ("source", "Source", 11, None),
    ("ref", "Ref", 11, None),
    ("category", "Category", 11, None),
    ("property_type", "Site type", 14, None),
    ("listing_type", "Sale/Rent", 9, None),
    ("title", "Title", 55, None),
    ("price_usd", "Price (USD)", 13, '"$"#,##0'),
    ("price_label", "Price as listed", 16, None),
    ("area_m2", "Area (m²)", 10, "#,##0"),
    ("price_per_m2", "USD / m²", 10, '"$"#,##0'),
    ("bedrooms", "Beds", 6, "0"),
    ("bathrooms", "Baths", 6, "0"),
    ("parking", "Parking", 8, "0"),
    ("location", "Location", 30, None),
    ("status", "Status", 10, None),
    ("furnished", "Furnished", 12, None),
    ("payment_type", "Payment", 10, None),
    ("agent", "Agent", 18, None),
    ("agent_phone", "Phone", 15, None),
    ("publish_date", "Published", 11, None),
    ("views", "Views", 7, "#,##0"),
    ("also_listed_on", "Possibly also listed as", 30, None),
    ("url", "Link", 9, None),
    ("description", "Description", 80, None),
]

thin = Side(style="thin", color=RULE)


def _median(s):
    s = s.dropna()
    return float(s.median()) if len(s) else None


def summary_by_category(df):
    sale = df[df.listing_type == "Sale"]
    rows = []
    for cat in ["Land", "Villa", "Apartment", "Commercial", "Other"]:
        d = sale[sale.category == cat]
        if d.empty:
            continue
        rows.append({
            "Category": cat,
            "Listings": len(d),
            "Confidence": int((d.source == "Confidence").sum()),
            "JSK": int((d.source == "JSK").sum()),
            "Median price": _median(d.price_usd),
            "Average price": d.price_usd.mean(),
            "Min price": d.price_usd.min(),
            "Max price": d.price_usd.max(),
            "Median area (m²)": _median(d.area_m2),
            "Median USD/m²": _median(d.price_per_m2),
        })
    return pd.DataFrame(rows)


def land_by_source(df):
    land = df[(df.category == "Land") & (df.listing_type == "Sale")]
    rows = []
    for src, d in [("All", land)] + list(land.groupby("source")):
        rows.append({
            "Source": src,
            "Plots": len(d),
            "Median USD/m²": _median(d.price_per_m2),
            "Average USD/m²": d.price_per_m2.mean(),
            "Cheapest USD/m²": d.price_per_m2.min(),
            "Most expensive USD/m²": d.price_per_m2.max(),
            "Median plot (m²)": _median(d.area_m2),
            "Total area (m²)": d.area_m2.sum(),
        })
    return pd.DataFrame(rows)


BEDROOM_LABELS = ["Studio / 1", "2", "3", "4", "5+", "Not listed"]


def bedroom_band(beds):
    if pd.isna(beds):
        return "Not listed"
    if beds <= 1:
        return "Studio / 1"
    return "5+" if beds >= 5 else str(int(beds))


def apartments_by_bedrooms(df, listing_type="Sale"):
    apts = df[(df.category == "Apartment") & (df.listing_type == listing_type)]
    apts = apts.assign(band=apts.bedrooms.map(bedroom_band))
    rows = []
    for band in BEDROOM_LABELS:
        d = apts[apts.band == band]
        if d.empty:
            continue
        rows.append({
            "Bedrooms": band,
            "Listings": len(d),
            "Median price": _median(d.price_usd),
            "Median area (m²)": _median(d.area_m2),
            "Median USD/m²": _median(d.price_per_m2),
        })
    return pd.DataFrame(rows)


def land_by_size(df):
    land = df[(df.category == "Land") & (df.listing_type == "Sale")].copy()
    land["band"] = pd.cut(land.area_m2, SIZE_BANDS, labels=SIZE_LABELS, right=False)
    rows = []
    for band in SIZE_LABELS:
        d = land[land.band == band]
        rows.append({
            "Plot size": band,
            "Plots": len(d),
            "Median price": _median(d.price_usd),
            "Median USD/m²": _median(d.price_per_m2),
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- writing helpers

def _write_table(ws, top, left, frame, formats, title=None):
    """Write a small styled summary table; returns (first data row, last row)."""
    if title:
        ws.cell(top, left, title).font = Font(bold=True, size=12, color=INK)
        top += 1
    for j, col in enumerate(frame.columns):
        c = ws.cell(top, left + j, col)
        c.font = Font(bold=True, color=MUTED, size=9)
        c.border = Border(bottom=Side(style="medium", color=INK))
        c.alignment = Alignment(horizontal="left" if j == 0 else "right", wrap_text=True)
    for i, row in enumerate(frame.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            if isinstance(value, float) and pd.isna(value):
                value = None
            c = ws.cell(top + i, left + j, value)
            c.border = Border(bottom=thin)
            c.alignment = Alignment(horizontal="left" if j == 0 else "right")
            fmt = formats.get(frame.columns[j])
            if fmt:
                c.number_format = fmt
    return top + 1, top + len(frame)


def _kpi(ws, row, col, label, value, fmt):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    lab = ws.cell(row, col, label)
    lab.font = Font(size=9, color=MUTED)
    val = ws.cell(row + 1, col, value)
    val.font = Font(size=18, bold=True, color=INK)
    val.number_format = fmt
    for r in (row, row + 1):
        for c in (col, col + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=TILE)
            ws.cell(r, c).alignment = Alignment(horizontal="left", vertical="center", indent=1)


def _write_data_sheet(wb, name, frame):
    ws = wb.create_sheet(name)
    cols = [c for c in DATA_COLUMNS if c[0] in frame.columns]
    for j, (_, header, width, _) in enumerate(cols, start=1):
        ws.cell(1, j, header)
        ws.column_dimensions[get_column_letter(j)].width = width
    for i, rec in enumerate(frame.to_dict("records"), start=2):
        for j, (key, _, _, fmt) in enumerate(cols, start=1):
            value = rec.get(key)
            if isinstance(value, float) and pd.isna(value):
                value = None
            if key == "description" and isinstance(value, str):
                value = value[:32000]  # Excel cell limit
            c = ws.cell(i, j, value)
            if fmt:
                c.number_format = fmt
            if key == "url" and value:
                c.value, c.hyperlink = "Open ↗", value
                c.font = Font(color="2A78D6", underline="single")
            if key in ("title", "description"):
                c.alignment = Alignment(vertical="top", wrap_text=key == "title")
    last = max(2, len(frame) + 1)
    ref = f"A1:{get_column_letter(len(cols))}{last}"
    table = Table(displayName=name.replace(" ", "_").replace("&", "and").replace("(", "").replace(")", ""), ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "C2"
    return ws


def _style_series(series, color, marker=False):
    if marker:
        series.marker.symbol = "circle"
        series.marker.size = 8
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = "FFFFFF"
        series.graphicalProperties.line.noFill = True
    else:
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color


# ---------------------------------------------------------------- main entry

def build_workbook(df: pd.DataFrame, path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    for col in range(1, 14):
        ws.column_dimensions[get_column_letter(col)].width = 14

    sale = df[df.listing_type == "Sale"]
    land = sale[sale.category == "Land"]
    villas = sale[sale.category == "Villa"]
    apts = sale[sale.category == "Apartment"]
    apts_rent = df[(df.listing_type == "Rent") & (df.category == "Apartment")]

    ws["A1"] = "Adma Real Estate — Land, Villas & Apartments Market Snapshot"
    ws["A1"].font = Font(size=18, bold=True, color=INK)
    ws["A2"] = (f"Scraped {datetime.now():%d %b %Y %H:%M} from confidencerealestate.com and jskre.com · "
                f"{len(df)} listings in Adma (Keserwan) · prices in USD")
    ws["A2"].font = Font(size=10, color=MUTED)

    kpis = [
        ("Adma listings (all)", len(df), "#,##0"),
        ("Land plots for sale", len(land), "#,##0"),
        ("Median land USD/m²", _median(land.price_per_m2), '"$"#,##0'),
        ("Median land price", _median(land.price_usd), '"$"#,##0'),
        ("Villas/houses for sale", len(villas), "#,##0"),
        ("Median villa price", _median(villas.price_usd), '"$"#,##0'),
    ]
    apt_kpis = [
        ("Apartments for sale", len(apts), "#,##0"),
        ("Median apartment price", _median(apts.price_usd), '"$"#,##0'),
        ("Median apartment USD/m²", _median(apts.price_per_m2), '"$"#,##0'),
        ("Median apartment size", _median(apts.area_m2), '#,##0" m²"'),
        ("Median bedrooms", _median(apts.bedrooms), "0"),
        ("Apartments for rent", len(apts_rent), "#,##0"),
    ]
    for k, (label, value, fmt) in enumerate(kpis):
        _kpi(ws, 4, 1 + 2 * k, label, value, fmt)
    for k, (label, value, fmt) in enumerate(apt_kpis):
        _kpi(ws, 7, 1 + 2 * k, label, value, fmt)

    money = '"$"#,##0'
    cat = summary_by_category(df)
    cat_first, cat_last = _write_table(
        ws, 11, 1, cat,
        {c: money for c in ["Median price", "Average price", "Min price", "Max price", "Median USD/m²"]}
        | {"Median area (m²)": "#,##0"},
        title="For-sale listings by category")

    src = land_by_source(df)
    src_top = cat_last + 3
    _, src_last = _write_table(
        ws, src_top, 1, src,
        {c: money for c in ["Median USD/m²", "Average USD/m²", "Cheapest USD/m²", "Most expensive USD/m²"]}
        | {"Median plot (m²)": "#,##0", "Total area (m²)": "#,##0"},
        title="Land for sale — price per m² by agency")

    size = land_by_size(df)
    size_top = src_last + 3
    size_first, size_last = _write_table(
        ws, size_top, 1, size, {"Median price": money, "Median USD/m²": money},
        title="Land for sale — by plot size")

    beds = apartments_by_bedrooms(df)
    beds_first, beds_last = _write_table(
        ws, size_last + 3, 1, beds,
        {"Median price": money, "Median USD/m²": money, "Median area (m²)": "#,##0"},
        title="Apartments for sale — by bedrooms")

    # Chart 1: for-sale listings per category, one series per agency.
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", "clustered", -10
    ch.title = "For-sale listings by category"
    for idx, source in ((3, "Confidence"), (4, "JSK")):
        s = Series(Reference(ws, min_col=idx, min_row=cat_first, max_row=cat_last), title=source)
        _style_series(s, SOURCE_COLORS[source])
        ch.series.append(s)
    ch.set_categories(Reference(ws, min_col=1, min_row=cat_first, max_row=cat_last))
    ch.y_axis.majorGridlines.spPr = None
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.height, ch.width = 7.5, 15
    ws.add_chart(ch, "L11")

    # Chart 2: median land USD/m² by plot size.
    ch2 = BarChart()
    ch2.type, ch2.title = "col", "Land: median USD/m² by plot size"
    s = Series(Reference(ws, min_col=4, min_row=size_first, max_row=size_last), title="Median USD/m²")
    _style_series(s, "2A78D6")
    ch2.series.append(s)
    ch2.set_categories(Reference(ws, min_col=1, min_row=size_first, max_row=size_last))
    ch2.legend = None
    ch2.x_axis.delete = ch2.y_axis.delete = False
    ch2.y_axis.number_format = '"$"#,##0'
    ch2.height, ch2.width = 7.5, 15
    ws.add_chart(ch2, "L27")

    # Chart 2b: apartments, median price by bedrooms.
    ch4 = BarChart()
    ch4.type, ch4.title = "col", "Apartments for sale: median price by bedrooms"
    # "Not listed" mixes unknown layouts (often whole floors) and would dwarf the other bars - table only.
    chart_last = beds_last - 1 if ws.cell(beds_last, 1).value == "Not listed" else beds_last
    s = Series(Reference(ws, min_col=3, min_row=beds_first, max_row=chart_last), title="Median price")
    _style_series(s, "1BAF7A")
    ch4.series.append(s)
    ch4.set_categories(Reference(ws, min_col=1, min_row=beds_first, max_row=chart_last))
    ch4.legend = None
    ch4.x_axis.delete = ch4.y_axis.delete = False
    ch4.y_axis.number_format = '"$"#,##0'
    ch4.height, ch4.width = 7.5, 15
    ws.add_chart(ch4, "L63")

    # Chart 3: land area vs price scatter (log axes - plots range from ~1,000 to 14,500 m²).
    land_sorted = df[df.category == "Land"].sort_values(["listing_type", "price_per_m2"], ascending=[False, True])
    _write_data_sheet(wb, "Land", land_sorted)
    ch3 = ScatterChart()
    ch3.title = "Land for sale: plot size vs asking price (log scale)"
    ch3.x_axis.title = "Plot area (m²)"
    ch3.x_axis.scaling.logBase = ch3.y_axis.scaling.logBase = 10
    ch3.x_axis.scaling.min, ch3.x_axis.scaling.max = 500, 20000
    ch3.y_axis.scaling.min, ch3.y_axis.scaling.max = 200000, 50000000
    ch3.x_axis.majorGridlines = None
    ch3.x_axis.number_format = "#,##0"
    ch3.y_axis.number_format = '"$"#,##0'
    ch3.x_axis.delete = False
    ch3.y_axis.delete = False
    # Excel scatter series need contiguous ranges, so stage each agency's for-sale plots on a hidden sheet.
    helper = wb.create_sheet("_chart_data")
    helper.sheet_state = "hidden"
    r = 1
    for source in ("Confidence", "JSK"):
        d = land_sorted[(land_sorted.source == source) & (land_sorted.listing_type == "Sale")]
        start = r
        for rec in d.itertuples():
            helper.cell(r, 1, rec.area_m2)
            helper.cell(r, 2, rec.price_usd)
            r += 1
        if r > start:
            xs = Reference(helper, min_col=1, min_row=start, max_row=r - 1)
            ys = Reference(helper, min_col=2, min_row=start, max_row=r - 1)
            ser = Series(ys, xs, title=source)
            _style_series(ser, SOURCE_COLORS[source], marker=True)
            ch3.series.append(ser)
    ch3.height, ch3.width = 9, 15
    ws.add_chart(ch3, "L43")

    _write_data_sheet(wb, "Villas", df[df.category == "Villa"].sort_values(["listing_type", "price_usd"], ascending=[False, True]))
    _write_data_sheet(wb, "Apartments", df[df.category == "Apartment"].sort_values(["listing_type", "price_per_m2"], ascending=[False, True]))
    _write_data_sheet(wb, "All Adma Listings", df.sort_values(["category", "listing_type", "price_usd"]))

    wb.move_sheet("_chart_data", offset=len(wb.sheetnames))
    wb.save(path)
    return path
